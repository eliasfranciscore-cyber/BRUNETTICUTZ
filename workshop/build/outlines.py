#!/usr/bin/env python3
"""
Genera el guion del facilitador a partir de los .pptx ya construidos.

Se hace así — y no a mano — para que el guion no se desincronice del deck: si cambia
una slide, se corre esto y el guion queda igual de actualizado. Las notas de orador
viven en el generador (`s.addNotes(...)`), que es la única fuente de verdad.

    python3 build/outlines.py
"""

from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
DECKS = ROOT / "decks"
OUT = DECKS / "outlines"

# (archivo del deck, salida, título, subtítulo)
VERSIONES = [
    ("ASCENSION-completo.pptx", "ascension-completo.md",
     "ASCENSIÓN — El barbero que cobra lo que vale",
     "La versión completa para la jornada de 6 horas: junta los tres borradores en un solo "
     "deck, con la paleta ASCENSIÓN, el logo real y un ejercicio por módulo."),
    ("V1-ascension-el-barbero-invisible.pptx", "v1-ascension.md",
     "V1 — ASCENSIÓN II · El barbero invisible",
     "El deck narrativo. Percepción → visibilidad → sistema. Pocas palabras por slide, "
     "una frase martillo por bloque."),
    ("V2-del-scroll-a-la-silla.pptx", "v2-del-scroll-a-la-silla.md",
     "V2 — Del scroll a la silla",
     "El deck del dinero. Abre con la economía cruda y cierra con la máquina de referidos."),
    ("V3-quince-segundos-el-taller.pptx", "v3-quince-segundos.md",
     "V3 — 15 segundos · el taller",
     "El deck laboratorio. Cada bloque termina en un ejercicio con cronómetro y la sala "
     "construye un video a lo largo del día."),
]


def altura_de(run):
    """Tamaño de fuente del run, o 0 si lo hereda del layout."""
    return run.font.size.pt if run.font.size else 0


def lineas_de(slide):
    """Todas las líneas de texto de la slide, con su tamaño de fuente."""
    out = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            texto = "".join(r.text for r in para.runs).strip()
            if not texto:
                continue
            tam = max((altura_de(r) for r in para.runs), default=0)
            out.append((tam, texto))
    return out


def titulo_de(lineas, etiqueta):
    """
    El texto más grande de la slide. En las slides de titular eso ES el titular; en las
    de lista (agenda, escala, criterios) el texto más grande se repite en cada fila, así
    que en ese caso el nombre de la slide sale del eyebrow.
    """
    if not lineas:
        return "(sin texto)"
    tam_max = max(t for t, _ in lineas)
    partes = [txt for t, txt in lineas if t == tam_max]
    if len(partes) > 2:
        return etiqueta.capitalize() if etiqueta else partes[0]
    return " / ".join(partes)[:120]


def etiqueta_de(lineas):
    """El eyebrow: la línea chica en mayúsculas que marca el módulo."""
    for tam, txt in lineas:
        if tam and tam <= 12 and txt == txt.upper() and len(txt) > 3:
            return txt
    return ""


def cuerpo_de(lineas, titulo, etiqueta):
    """El resto del texto, en orden, sin repetir el titular ni el eyebrow."""
    vistos = set(titulo.split(" / ")) | {etiqueta}
    fuera = {"BRUNETTICUTZ", "ASCENSIÓN", "DEL SCROLL A LA SILLA", "15 SEGUNDOS"}
    out = []
    for _, txt in lineas:
        if txt in vistos or txt in fuera:
            continue
        vistos.add(txt)
        out.append(" ".join(txt.split()))
    return out


def escribir(deck_file, out_file, titulo, subtitulo):
    pres = Presentation(DECKS / deck_file)
    total = len(pres.slides)

    md = [
        f"# {titulo}",
        "",
        subtitulo,
        "",
        f"**{total} slides.** Generado desde `decks/{deck_file}` con `build/outlines.py` — "
        "no editar a mano: los cambios se hacen en el generador del deck y se vuelve a correr.",
        "",
        "Las notas de orador de este documento son las mismas que están dentro del `.pptx`, "
        "así que se pueden leer desde la vista de presentador de PowerPoint o Keynote.",
        "",
        "---",
        "",
    ]

    for i, slide in enumerate(pres.slides, start=1):
        lineas = lineas_de(slide)
        etiqueta = etiqueta_de(lineas)
        titular = titulo_de(lineas, etiqueta)
        cuerpo = cuerpo_de(lineas, titular, etiqueta)

        cab = f"## {i:02d} · {titular}"
        md.append(cab)
        if etiqueta:
            md.append(f"`{etiqueta}`")
        md.append("")

        if cuerpo:
            md.append("**En pantalla**")
            md.append("")
            for c in cuerpo[:14]:
                md.append(f"- {c}")
            if len(cuerpo) > 14:
                md.append(f"- *(+{len(cuerpo) - 14} elementos más en la slide)*")
            md.append("")

        notas = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if notas:
            md.append("**Qué dices**")
            md.append("")
            for parrafo in [p.strip() for p in notas.split("\n") if p.strip()]:
                md.append(f"> {parrafo}")
                md.append(">")
            md.pop()
            md.append("")

        md.append("---")
        md.append("")

    (OUT / out_file).write_text("\n".join(md), encoding="utf-8")
    print(f"✔ decks/outlines/{out_file}  ({total} slides)")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for args in VERSIONES:
        escribir(*args)
